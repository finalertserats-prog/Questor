#!/usr/bin/env python3
"""
Turn docs/QUESTOR-REFERENCE.md into a Word file and a print-ready HTML page.

The reference is the document somebody takes over the codebase from, so it has
to leave the repository in a form people outside it can open. Markdown is the
source of truth — it is versioned, diffable and reviewed with the code — and
these are renderings of it, regenerated rather than edited.

The HTML is styled for print: the PDF is produced from it with

    chrome --headless --disable-gpu --print-to-pdf=<out> --no-pdf-header-footer <html>

rather than by a Python PDF library, because the document is 43,000 words of
headings, tables and code, and a browser's layout engine handles page breaks,
table widths and monospace runs correctly where a hand-rolled writer does not.
"""

from __future__ import annotations

import html
import re
import sys
from pathlib import Path

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

REPO = Path(__file__).resolve().parent.parent
SOURCE = REPO / "docs" / "QUESTOR-REFERENCE.md"
OUT = REPO / "deliverables"

INK = RGBColor(0x13, 0x1A, 0x26)
QUIET = RGBColor(0x4A, 0x54, 0x68)


def read_source() -> str:
    if not SOURCE.exists():
        sys.exit(f"Not found: {SOURCE}")
    return SOURCE.read_text(encoding="utf-8")


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

def _runs(paragraph, text: str) -> None:
    """Inline bold, italic and code, so emphasis survives the conversion."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def to_word(md: str, out: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK

    lines = md.split("\n")
    i = 0
    in_code = False
    code: list[str] = []

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("```"):
            if in_code:
                para = doc.add_paragraph()
                run = para.add_run("\n".join(code))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = QUIET
                code = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code.append(line)
            i += 1
            continue

        # A table: a header row, a rule, then body rows.
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[\s:|-]+$", "|".join(cells)):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=width)
                table.style = "Light Grid Accent 1"
                for r, cells in enumerate(rows):
                    row = table.add_row().cells
                    for c in range(width):
                        text = re.sub(r"[*`]", "", cells[c] if c < len(cells) else "")
                        row[c].text = text
                        if r == 0:
                            for para in row[c].paragraphs:
                                for run in para.runs:
                                    run.bold = True
                doc.add_paragraph()
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(re.sub(r"[*`]", "", line[level:].strip()), level=min(level, 4))
        elif re.match(r"^\s*[-*]\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            _runs(para, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            para = doc.add_paragraph(style="List Number")
            _runs(para, re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.startswith(">"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(24)
            _runs(para, line.lstrip("> "))
            for run in para.runs:
                run.italic = True
                run.font.color.rgb = QUIET
        elif line.strip() in {"---", "***", "___"}:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run("· · ·").font.color.rgb = QUIET
        elif line.strip():
            _runs(doc.add_paragraph(), line)

        i += 1

    doc.save(out)


# --------------------------------------------------------------------------
# Print-ready HTML, which Chrome turns into the PDF
# --------------------------------------------------------------------------

CSS = """
@page { size: A4; margin: 18mm 16mm; }
body { font-family: Georgia, 'Times New Roman', serif; font-size: 10.5pt;
       line-height: 1.55; color: #131a26; max-width: none; margin: 0; }
h1 { font-size: 24pt; margin: 0 0 6pt; page-break-before: always; line-height: 1.15; }
h1:first-of-type { page-break-before: avoid; }
h2 { font-size: 16pt; margin: 18pt 0 6pt; border-bottom: 1px solid #d2d7dd; padding-bottom: 3pt; }
h3 { font-size: 12.5pt; margin: 14pt 0 4pt; }
h4 { font-size: 11pt; margin: 12pt 0 4pt; color: #4a5468; }
h1, h2, h3, h4 { page-break-after: avoid; font-family: 'Segoe UI', Calibri, sans-serif; }
p, li { orphans: 3; widows: 3; }
code { font-family: Consolas, 'Courier New', monospace; font-size: 9pt;
       background: #f1f3f5; padding: 0 3px; border-radius: 2px; }
pre { font-family: Consolas, monospace; font-size: 8.5pt; background: #f7f8f9;
      border: 1px solid #e5e8ec; border-radius: 3px; padding: 8pt;
      white-space: pre-wrap; word-wrap: break-word; page-break-inside: avoid; }
pre code { background: none; padding: 0; }
table { border-collapse: collapse; width: 100%; font-size: 9pt; margin: 8pt 0;
        page-break-inside: avoid; }
th, td { border: 1px solid #d2d7dd; padding: 4pt 6pt; text-align: left;
         vertical-align: top; }
th { background: #eff1f3; font-weight: 600; }
blockquote { margin: 8pt 0 8pt 12pt; padding-left: 10pt; border-left: 2px solid #b8c0c6;
             color: #4a5468; font-style: italic; }
hr { border: 0; border-top: 1px solid #d2d7dd; margin: 14pt 0; }
a { color: #131a26; text-decoration: none; }
"""


def to_html(md: str, out: Path) -> None:
    body = markdown.markdown(
        md, extensions=["tables", "fenced_code", "toc", "sane_lists", "attr_list"]
    )
    out.write_text(
        "<!doctype html>\n<html lang='en'><head><meta charset='utf-8'>\n"
        "<title>Questor — Reference</title>\n"
        f"<style>{CSS}</style>\n</head><body>\n{body}\n</body></html>\n",
        encoding="utf-8",
    )


def main() -> None:
    md = read_source()
    OUT.mkdir(exist_ok=True)
    docx_path = OUT / "Questor-Reference.docx"
    html_path = OUT / "Questor-Reference.html"
    to_word(md, docx_path)
    to_html(md, html_path)
    print(f"word  {docx_path}  ({docx_path.stat().st_size // 1024} KB)")
    print(f"html  {html_path}  ({html_path.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
